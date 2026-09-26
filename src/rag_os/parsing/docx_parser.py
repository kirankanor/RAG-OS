from __future__ import annotations

from pathlib import Path

from rag_os.core.types import Document
from rag_os.parsing.base import Parser, parser_registry


@parser_registry.register(
    "docx_python_docx",
    "Extracts paragraph text (and simple tables) from .docx files via python-docx.",
)
class DocxParser(Parser):
    name = "docx_python_docx"

    @classmethod
    def supported_extensions(cls) -> tuple[str, ...]:
        return (".docx",)

    def parse(self, file_path: str | Path) -> Document:
        try:
            import docx
        except ImportError as e:
            raise ImportError(
                "python-docx is not installed. Run: uv add python-docx"
            ) from e

        path = Path(file_path)
        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]

        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))

        text = "\n".join(parts)
        return Document(
            source_filename=path.name,
            text=text,
            parser_name=self.name,
            metadata={"num_paragraphs": len(doc.paragraphs), "num_chars": len(text)},
        )
